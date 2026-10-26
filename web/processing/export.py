"""
export.py — XLSX export, ported from client/export.py with the same sheets,
headers and rules, reading cells from the database instead of sidecar files:

- one Summary row per requested sample, one Cells row per kept cell;
- a channel that isn't processed (or doesn't exist for the sample) is `###`
  (NO_DATA) in the summary and contributes no cell rows, while a processed
  channel with zero kept cells is a real 0 — "we don't know yet" and "we counted
  zero" stay distinguishable;
- coexpression comes from processing.coexpression, the same code the Composite
  tab uses, so the two can never disagree;
- the rate is ### when pairs or the SNAP population are unknown, and when the
  population is zero (a rate against zero cells is undefined, not zero).
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from processing.coexpression import compute_coexpression
from projects.models import CHANNELS, CellStatus, ImageStatus, Project

NO_DATA = "###"

SUMMARY_HEADERS = ["prefix", "snap_kept", "cck_kept", "chr_kept", "coexpressing_pairs",
                   "coexpression_rate_pct"]
CELL_HEADERS = ["prefix", "channel", "cell_id", "status", "area", "centroid_x", "centroid_y",
                "source", "edited", "coexpressing"]


def _kept_cells(image) -> list[dict]:
    return [c.as_dict() for c in image.cells.filter(status=CellStatus.KEPT).order_by("cell_id")]


def build_export_rows(project: Project, prefixes: list[str]) -> tuple[list[dict], list[dict]]:
    samples = {s.prefix: s for s in project.samples.filter(prefix__in=prefixes)}
    summary_rows: list[dict] = []
    cell_rows: list[dict] = []

    for prefix in prefixes:
        sample = samples.get(prefix)
        images = {img.channel: img for img in sample.images.all()} if sample else {}
        kept_by_channel: dict[str, list[dict] | None] = {}
        for ch in CHANNELS:
            img = images.get(ch)
            kept_by_channel[ch] = _kept_cells(img) if img and img.status == ImageStatus.DONE else None

        snap_kept, cck_kept, chr_kept = (kept_by_channel[ch] for ch in CHANNELS)

        if cck_kept is not None and chr_kept is not None:
            result = compute_coexpression(cck_kept, chr_kept)
            n_pairs = len(result.pairs)
        else:
            result = None
            n_pairs = None

        if n_pairs is not None and snap_kept is not None and len(snap_kept) > 0:
            rate = round(n_pairs / len(snap_kept) * 100, 2)
        else:
            rate = None

        summary_rows.append({
            "prefix": prefix,
            "snap_kept": len(snap_kept) if snap_kept is not None else NO_DATA,
            "cck_kept": len(cck_kept) if cck_kept is not None else NO_DATA,
            "chr_kept": len(chr_kept) if chr_kept is not None else NO_DATA,
            "coexpressing_pairs": n_pairs if n_pairs is not None else NO_DATA,
            "coexpression_rate_pct": rate if rate is not None else NO_DATA,
        })

        for channel, cells, coexpr_ids in (
            ("SNAP", snap_kept, None),
            ("CCK", cck_kept, result.cck_ids if result else None),
            ("CHR", chr_kept, result.chr_ids if result else None),
        ):
            if cells is None:
                continue
            for c in cells:
                if channel == "SNAP":
                    coexpressing = ""  # not applicable: SNAP isn't part of the CCK/CHR overlap
                elif coexpr_ids is not None:
                    coexpressing = c["id"] in coexpr_ids
                else:
                    coexpressing = NO_DATA  # this channel has data but its CCK/CHR partner doesn't
                cell_rows.append({
                    "prefix": prefix,
                    "channel": channel,
                    "cell_id": c["id"],
                    "status": c["status"],
                    "area": c["area"],
                    "centroid_x": c["centroid"][0],
                    "centroid_y": c["centroid"][1],
                    "source": c["source"],
                    "edited": c.get("edited", False),
                    "coexpressing": coexpressing,
                })

    return summary_rows, cell_rows


def export_xlsx(path: Path, summary_rows: list[dict] | None, cell_rows: list[dict] | None) -> Path:
    """A sheet whose rows are None is left out entirely."""
    wb = openpyxl.Workbook(write_only=True)
    if summary_rows is not None:
        ws = wb.create_sheet("Summary")
        ws.append(SUMMARY_HEADERS)
        for row in summary_rows:
            ws.append([row[h] for h in SUMMARY_HEADERS])
    if cell_rows is not None:
        ws = wb.create_sheet("Cells")
        ws.append(CELL_HEADERS)
        for row in cell_rows:
            ws.append([row[h] for h in CELL_HEADERS])
    tmp = path.with_name(path.name + ".tmp")
    wb.save(tmp)
    tmp.replace(path)
    return path


def download_filename(project_name: str) -> str:
    safe = "".join("_" if ch in '/\\:*?"<>|' or ord(ch) < 32 else ch for ch in project_name).strip() or "project"
    return f"{safe}_export.xlsx"
